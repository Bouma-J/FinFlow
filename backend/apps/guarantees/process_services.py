"""Services métier : main levée et dation en paiement."""
from __future__ import annotations

from datetime import date
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from apps.corebanking.services import (
    CoreBankingError,
    assert_client_outstanding_for_dation,
    assert_loan_settled,
    get_client_outstanding,
    get_loan_status,
)
from apps.credits.amounts import reference_amount
from apps.tenants.currency import tenant_currency
from apps.workflow.models import WorkflowDefinition
from apps.workflow.services import WorkflowError, start_workflow

from django.core.files.base import ContentFile

from .models import (
    DationAsset,
    DationFee,
    DationRequest,
    Guarantee,
    GuaranteeMovement,
    GuaranteeReleaseRequest,
    ReleaseFee,
)

_OPEN_DATION_STATUSES = (
    DationRequest.Status.DRAFT,
    DationRequest.Status.IN_APPROVAL,
    DationRequest.Status.RETURNED,
    DationRequest.Status.APPROVED,
    DationRequest.Status.BLOCKED,
)

_OPEN_RELEASE_STATUSES = (
    GuaranteeReleaseRequest.Status.DRAFT,
    GuaranteeReleaseRequest.Status.IN_APPROVAL,
    GuaranteeReleaseRequest.Status.RETURNED,
    GuaranteeReleaseRequest.Status.APPROVED,
    GuaranteeReleaseRequest.Status.BLOCKED,
)


class ProcessError(Exception):
    """Erreur métier des processus main levée / dation."""


def _next_reference(prefix: str, model, tenant_id) -> str:
    today = timezone.now().strftime("%Y%m%d")
    count = (
        model.all_tenants.filter(tenant_id=tenant_id, created_at__date=date.today())
        .count()
        + 1
    )
    return f"{prefix}-{today}-{count:04d}"


def _loan_cbs_ref(loan) -> str:
    """Préfère refDemande Perfect, puis références contrat / legacy."""
    if loan is None:
        return ""
    return (
        (getattr(loan, "cbs_demande_ref", None) or "").strip()
        or (getattr(loan, "core_banking_reference", None) or "").strip()
        or (getattr(loan, "cbs_contract_number", None) or "").strip()
        or (getattr(loan, "cbs_demande_number", None) or "").strip()
    )


def _client_cbs_ids(client) -> dict:
    if client is None:
        return {}
    return {
        "code_adherent": (getattr(client, "cbs_client_id", None) or "").strip(),
        "num_manuel": (getattr(client, "cbs_account_number", None) or "").strip(),
        "num_piece_identite": (getattr(client, "national_id", None) or "").strip(),
    }


def _resolve_loan_ref(guarantee, loan=None, cbs_loan_reference="") -> tuple:
    """Retourne (loan, cbs_ref) — refDemande Perfect prioritaire."""
    ref = (cbs_loan_reference or "").strip()
    resolved_loan = loan
    if resolved_loan is None and guarantee.application_id:
        resolved_loan = getattr(guarantee.application, "loan", None)
        if resolved_loan is None:
            from apps.credits.models import Loan

            resolved_loan = Loan.objects.filter(
                application_id=guarantee.application_id
            ).first()
    if not ref:
        ref = _loan_cbs_ref(resolved_loan)
    return resolved_loan, ref


def _assert_release_editable(request: GuaranteeReleaseRequest):
    if request.status not in (
        GuaranteeReleaseRequest.Status.DRAFT,
        GuaranteeReleaseRequest.Status.RETURNED,
    ):
        raise ProcessError(
            "Seuls les dossiers brouillon ou retournés peuvent être modifiés."
        )


def _parse_release_date(request_date):
    if isinstance(request_date, str) and request_date.strip():
        try:
            return date.fromisoformat(request_date.strip()[:10])
        except ValueError as exc:
            raise ProcessError("Date de demande invalide.") from exc
    if request_date:
        return request_date
    return timezone.now().date()


def _create_release_fee_lines(request: GuaranteeReleaseRequest, fees):
    if not fees:
        return
    if not isinstance(fees, list):
        raise ProcessError("fees doit être une liste.")
    for item in fees:
        if not isinstance(item, dict):
            raise ProcessError("Chaque frais doit être un objet.")
        amount = _to_decimal_or_none(item.get("amount"))
        if amount is None:
            raise ProcessError("Le montant du frais est obligatoire.")
        fee_type = (item.get("fee_type") or ReleaseFee.FeeType.OTHER).strip()
        if fee_type not in ReleaseFee.FeeType.values:
            fee_type = ReleaseFee.FeeType.OTHER
        payer = (item.get("payer") or ReleaseFee.Payer.CLIENT).strip()
        if payer not in ReleaseFee.Payer.values:
            payer = ReleaseFee.Payer.CLIENT
        label = (item.get("label") or "").strip()
        if not label:
            label = dict(ReleaseFee.FeeType.choices).get(fee_type, "Frais")
        ReleaseFee.objects.create(
            tenant_id=request.tenant_id,
            release=request,
            fee_type=fee_type,
            label=label,
            amount=amount,
            payer=payer,
            fee_date=item.get("fee_date") or None,
            recoverable=bool(item.get("recoverable", True)),
            notes=(item.get("notes") or "").strip(),
        )


def ensure_release_document_categories(tenant) -> int:
    from apps.documents.category_seed import ensure_document_categories

    return ensure_document_categories(
        tenant,
        (
            ("ML_DEMANDE", "Demande de main levée (client)"),
            ("ML_ACTE_SIGNE", "Acte de main levée signé"),
            ("ML_QUITTANCE", "Quittance / preuve de solde"),
            ("ML_RADIATION", "Preuve de radiation"),
            ("ML_OTHER", "Autre pièce main levée"),
        ),
    )


def release_documents(request: GuaranteeReleaseRequest):
    from django.contrib.contenttypes.models import ContentType

    from apps.documents.models import Document

    ct = ContentType.objects.get_for_model(GuaranteeReleaseRequest)
    return Document.objects.filter(
        content_type=ct, object_id=request.id
    ).select_related("category", "uploaded_by")


def has_client_demande(request: GuaranteeReleaseRequest) -> bool:
    return release_documents(request).filter(category__code="ML_DEMANDE").exists()


@transaction.atomic
def initiate_release_request(
    *,
    guarantee,
    user,
    comment="",
    cbs_loan_reference="",
    loan=None,
    request_date=None,
    release_fees=None,
    cbs_client_id="",
    fees=None,
    as_draft=True,
):
    """
    Vérifie CBS (prêt soldé) puis crée la demande.

    Par défaut en brouillon pour permettre le dépôt de la demande client,
    la génération de l'acte et les frais avant soumission.
    """
    if guarantee.status != Guarantee.Status.ACTIVE:
        raise ProcessError(
            "Seule une garantie active peut faire l'objet d'une main levée."
        )
    if GuaranteeReleaseRequest.objects.filter(
        guarantee=guarantee,
        status__in=_OPEN_RELEASE_STATUSES,
    ).exists():
        raise ProcessError(
            "Une demande de main levée est déjà en cours pour cette garantie."
        )

    loan_obj, loan_ref = _resolve_loan_ref(guarantee, loan, cbs_loan_reference)
    if not loan_ref:
        raise ProcessError(
            "Référence demande crédit CBS (refDemande) manquante : "
            "composez la main levée depuis Mains levées → Nouvelle et "
            "sélectionnez un crédit soldé."
        )
    currency = tenant_currency(guarantee.tenant_id)
    if guarantee.application_id and getattr(guarantee.application, "currency", None):
        currency = guarantee.application.currency

    client_ids = _client_cbs_ids(guarantee.client)
    if cbs_client_id:
        client_ids["code_adherent"] = cbs_client_id.strip()

    try:
        cbs = assert_loan_settled(
            guarantee.tenant_id,
            loan_ref,
            currency=currency,
            **client_ids,
        )
    except CoreBankingError as exc:
        raise ProcessError(str(exc)) from exc

    fees_amount = None
    if release_fees not in (None, ""):
        try:
            fees_amount = Decimal(str(release_fees))
        except (InvalidOperation, TypeError) as exc:
            raise ProcessError("Frais de main levée invalides.") from exc

    parsed_date = _parse_release_date(request_date)
    cbs_id = (
        client_ids.get("code_adherent")
        or getattr(guarantee.client, "cbs_client_id", "")
        or ""
    ).strip()

    req = GuaranteeReleaseRequest(
        guarantee=guarantee,
        application_id=guarantee.application_id
        or getattr(loan_obj, "application_id", None),
        loan=loan_obj,
        agency_id=guarantee.agency_id,
        cbs_loan_reference=loan_ref,
        cbs_client_id=cbs_id,
        cbs_settled=cbs["settled"],
        cbs_outstanding=cbs["outstanding"],
        cbs_currency=cbs["currency"],
        cbs_checked_at=timezone.now(),
        cbs_raw=cbs.get("raw") or {},
        request_date=parsed_date,
        release_fees=fees_amount,
        status=GuaranteeReleaseRequest.Status.DRAFT,
        comment=comment or "",
        created_by=user,
        updated_by=user,
    )
    req.reference = _next_reference("ML", GuaranteeReleaseRequest, guarantee.tenant_id)
    req.save()
    ensure_release_document_categories(guarantee.tenant)
    _create_release_fee_lines(req, fees or [])
    req.apply_fees_snapshot(persist=True)

    if as_draft:
        return req
    return submit_release_request(req, user=user)


@transaction.atomic
def update_release_request(request: GuaranteeReleaseRequest, *, user=None, **fields):
    _assert_release_editable(request)
    if "comment" in fields and fields["comment"] is not None:
        request.comment = fields["comment"]
    if "request_date" in fields and fields["request_date"] is not None:
        request.request_date = _parse_release_date(fields["request_date"])
    if "cbs_loan_reference" in fields and fields["cbs_loan_reference"] is not None:
        request.cbs_loan_reference = (fields["cbs_loan_reference"] or "").strip()
    if "loan" in fields:
        request.loan = fields["loan"]
    if user is not None:
        request.updated_by = user
    request.save()
    return request


@transaction.atomic
def add_release_fee(request: GuaranteeReleaseRequest, data: dict, *, user=None):
    _assert_release_editable(request)
    _create_release_fee_lines(request, [data])
    if user is not None:
        request.updated_by = user
        request.save(update_fields=["updated_by", "updated_at"])
    request.apply_fees_snapshot(persist=True)
    return request.fees.order_by("-created_at").first()


@transaction.atomic
def remove_release_fee(request: GuaranteeReleaseRequest, fee_id, *, user=None):
    _assert_release_editable(request)
    try:
        fee = request.fees.get(pk=fee_id)
    except ReleaseFee.DoesNotExist as exc:
        raise ProcessError("Frais introuvable sur ce dossier.") from exc
    fee.delete()
    if user is not None:
        request.updated_by = user
        request.save(update_fields=["updated_by", "updated_at"])
    request.apply_fees_snapshot(persist=True)
    return request


@transaction.atomic
def refresh_release_cbs(request: GuaranteeReleaseRequest, *, user=None):
    if request.status in (
        GuaranteeReleaseRequest.Status.COMPLETED,
        GuaranteeReleaseRequest.Status.CANCELLED,
        GuaranteeReleaseRequest.Status.REJECTED,
    ):
        raise ProcessError("Impossible de rafraîchir le CBS sur ce statut.")
    currency = request.cbs_currency or tenant_currency(request.tenant_id)
    client = getattr(request.guarantee, "client", None)
    client_ids = _client_cbs_ids(client)
    if request.cbs_client_id:
        client_ids["code_adherent"] = request.cbs_client_id.strip()
    try:
        cbs = get_loan_status(
            request.tenant_id,
            request.cbs_loan_reference,
            currency=currency,
            **client_ids,
        )
    except CoreBankingError as exc:
        raise ProcessError(str(exc)) from exc
    request.cbs_settled = cbs["settled"]
    request.cbs_outstanding = cbs["outstanding"]
    request.cbs_currency = cbs["currency"]
    request.cbs_checked_at = timezone.now()
    request.cbs_raw = cbs.get("raw") or {}
    if user is not None:
        request.updated_by = user
    request.save()
    return request


@transaction.atomic
def submit_release_request(request: GuaranteeReleaseRequest, *, user=None):
    if request.status not in (
        GuaranteeReleaseRequest.Status.DRAFT,
        GuaranteeReleaseRequest.Status.RETURNED,
    ):
        raise ProcessError(
            "Seuls les dossiers brouillon ou retournés peuvent être soumis."
        )
    refresh_release_cbs(request, user=user)
    if not request.cbs_settled:
        raise ProcessError(
            "Main levée impossible : le prêt n'est pas soldé dans le Core Banking "
            f"(encours restant : {request.cbs_outstanding} "
            f"{request.cbs_currency or ''})."
        )
    if not has_client_demande(request):
        raise ProcessError(
            "La demande de main levée du client doit être jointe avant soumission."
        )
    if not request.has_generated_acte():
        raise ProcessError(
            "Générez l'acte de main levée avant de soumettre le dossier."
        )

    amount = (
        request.release_fees
        or request.fees_client_total
        or request.guarantee.current_value
        or request.guarantee.expertise_value
        or 0
    )
    try:
        start_workflow(
            request,
            amount=amount,
            risk_level=None,
            target_type=WorkflowDefinition.TargetType.MAIN_LEVEE,
        )
    except WorkflowError as exc:
        raise ProcessError(str(exc)) from exc

    request.status = GuaranteeReleaseRequest.Status.IN_APPROVAL
    if user is not None:
        request.updated_by = user
    request.save(
        update_fields=["status", "updated_by", "updated_at"]
        if user
        else ["status", "updated_at"]
    )
    return request


@transaction.atomic
def cancel_release_request(request: GuaranteeReleaseRequest, *, user=None, comment=""):
    if request.status not in (
        GuaranteeReleaseRequest.Status.DRAFT,
        GuaranteeReleaseRequest.Status.RETURNED,
        GuaranteeReleaseRequest.Status.IN_APPROVAL,
        GuaranteeReleaseRequest.Status.BLOCKED,
    ):
        raise ProcessError("Ce dossier ne peut plus être annulé.")
    from apps.workflow.services import cancel_active_workflows_for_target

    cancel_active_workflows_for_target(request)
    request.status = GuaranteeReleaseRequest.Status.CANCELLED
    if comment:
        request.comment = (
            (request.comment + "\n" if request.comment else "") + comment
        )
    if user is not None:
        request.updated_by = user
    request.save()
    return request


def _build_release_acte_docx(request: GuaranteeReleaseRequest) -> bytes:
    """Génère un acte DOCX minimal à partir des données du dossier."""
    from io import BytesIO

    from docx import Document

    guarantee = request.guarantee
    client = guarantee.client
    tenant = request.tenant
    doc = Document()
    doc.add_heading("ACTE DE MAIN LEVÉE", level=1)
    doc.add_paragraph(f"Référence : {request.reference}")
    doc.add_paragraph(f"Date : {(request.request_date or timezone.localdate()).isoformat()}")
    doc.add_paragraph("")
    doc.add_paragraph(f"Filiale : {getattr(tenant, 'name', '')}")
    doc.add_paragraph(
        f"Client : {getattr(client, 'display_name', str(client))}"
    )
    doc.add_paragraph(f"Matricule CBS client : {request.cbs_client_id or '—'}")
    doc.add_paragraph("")
    doc.add_paragraph(
        f"Garantie : {guarantee.reference or guarantee.pk} — "
        f"{guarantee.get_guarantee_type_display()}"
    )
    doc.add_paragraph(f"Description : {guarantee.description or '—'}")
    value = guarantee.current_value or guarantee.expertise_value
    doc.add_paragraph(
        f"Valeur : {value} {request.cbs_currency or tenant_currency(request.tenant_id)}"
        if value is not None
        else "Valeur : —"
    )
    doc.add_paragraph("")
    doc.add_paragraph(f"Référence prêt CBS : {request.cbs_loan_reference or '—'}")
    doc.add_paragraph(
        f"Encours CBS : {request.cbs_outstanding} {request.cbs_currency}"
        if request.cbs_outstanding is not None
        else "Encours CBS : —"
    )
    doc.add_paragraph(
        f"Prêt soldé (CBS) : {'Oui' if request.cbs_settled else 'Non'}"
    )
    fees = request.apply_fees_snapshot(persist=False)
    doc.add_paragraph(
        f"Frais client : {fees['fees_client_total'] or 0} | "
        f"Frais institution : {fees['fees_institution_total'] or 0}"
    )
    if request.comment:
        doc.add_paragraph("")
        doc.add_paragraph(f"Commentaire : {request.comment}")
    doc.add_paragraph("")
    doc.add_paragraph(
        "Par le présent acte, l'établissement consent la main levée "
        "de la garantie susvisée, le prêt ayant été constaté soldé "
        "auprès du Core Banking."
    )
    doc.add_paragraph("")
    doc.add_paragraph("Fait pour valoir ce que de droit.")
    doc.add_paragraph("")
    doc.add_paragraph("Signature de l'établissement : ____________________")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@transaction.atomic
def generate_release_acte(request: GuaranteeReleaseRequest, *, user=None):
    if request.status in (
        GuaranteeReleaseRequest.Status.COMPLETED,
        GuaranteeReleaseRequest.Status.CANCELLED,
        GuaranteeReleaseRequest.Status.REJECTED,
    ):
        raise ProcessError("Impossible de générer l'acte sur ce statut.")
    if request.acte_status == GuaranteeReleaseRequest.ActeStatus.SIGNED:
        raise ProcessError(
            "Un acte signé est déjà déposé ; annulez-le avant de régénérer."
        )
    content = _build_release_acte_docx(request)
    filename = f"acte-main-levee-{request.reference or request.pk}.docx"
    request.acte_generated.save(filename, ContentFile(content), save=False)
    request.acte_generated_at = timezone.now()
    request.acte_status = GuaranteeReleaseRequest.ActeStatus.GENERATED
    if user is not None:
        request.updated_by = user
    request.save()
    return request


def user_can_deposit_release_acte(user, request: GuaranteeReleaseRequest) -> bool:
    """
    Initiateur (droit Django) ou membre d'un groupe d'étape du circuit
    de la demande peut déposer l'acte signé / clôturer.
    """
    if not (user and user.is_authenticated):
        return False
    if user.is_superuser or getattr(user, "is_group_level", False):
        return True
    if user.has_perm("guarantees.initiate_guaranteereleaserequest"):
        return True
    from django.contrib.contenttypes.models import ContentType

    from apps.workflow.models import ApprovalTask, WorkflowInstance

    ct = ContentType.objects.get_for_model(GuaranteeReleaseRequest)
    inst = (
        WorkflowInstance.objects.filter(content_type=ct, object_id=request.pk)
        .order_by("-created_at")
        .first()
    )
    if inst is None:
        return False
    user_group_ids = set(user.groups.values_list("id", flat=True))
    if not user_group_ids:
        return False
    step_group_ids = set(
        ApprovalTask.objects.filter(instance=inst).values_list(
            "step__required_group_id", flat=True
        )
    )
    return bool(user_group_ids.intersection(step_group_ids))


@transaction.atomic
def upload_release_acte_signed(request: GuaranteeReleaseRequest, upload, *, user=None):
    """Dépose l'acte signé (obligatoire avant clôture — choix métier A)."""
    if request.status in (
        GuaranteeReleaseRequest.Status.CANCELLED,
        GuaranteeReleaseRequest.Status.REJECTED,
    ):
        raise ProcessError("Impossible de déposer l'acte signé sur ce statut.")
    if not request.has_generated_acte():
        raise ProcessError("Générez d'abord l'acte de main levée.")
    name = getattr(upload, "name", None) or "acte-signe.pdf"
    request.acte_signed.save(name, upload, save=False)
    request.acte_signed_at = timezone.now()
    request.acte_status = GuaranteeReleaseRequest.ActeStatus.SIGNED
    if user is not None:
        request.updated_by = user
    request.save()

    # Index GED pour consultation unifiée.
    from django.contrib.contenttypes.models import ContentType

    from apps.documents.models import Document, DocumentCategory
    from apps.documents.quotas import bump_ged_usage

    ensure_release_document_categories(request.tenant)
    category = DocumentCategory.objects.filter(
        tenant_id=request.tenant_id, code="ML_ACTE_SIGNE", is_active=True
    ).first()
    if category is not None:
        ct = ContentType.objects.get_for_model(GuaranteeReleaseRequest)
        # Relire le fichier depuis le FileField (le stream upload peut être consommé).
        doc = Document(
            tenant_id=request.tenant_id,
            category=category,
            name=f"Acte signé {request.reference}",
            content_type=ct,
            object_id=request.id,
            uploaded_by=user,
        )
        doc.mime_type = getattr(upload, "content_type", "") or ""
        if request.acte_signed:
            doc.file.save(
                name,
                ContentFile(request.acte_signed.read()),
                save=False,
            )
            request.acte_signed.seek(0)
        doc.save()
        doc.compute_hash()
        doc.save(update_fields=["mime_type", "sha256", "size_bytes"])
        bump_ged_usage(doc.tenant_id, doc.size_bytes)

    # Si déjà approuvée, clôturer automatiquement (acte signé obligatoire).
    if request.status == GuaranteeReleaseRequest.Status.APPROVED:
        request = complete_release_request(request)
    return request


def release_client_context(*, client, tenant_id=None) -> dict:
    """
    Contexte pour composer une main levée :
    - matricule CBS (fiche client) ;
    - garanties actives Fin Flow ;
    - crédits / prêts du client avec statut CBS (soldé / actif).
    """
    from apps.credits.models import CreditApplication, Loan
    from apps.guarantees.serializers import GuaranteeSerializer

    tid = tenant_id or client.tenant_id
    cbs_id = (getattr(client, "cbs_client_id", "") or "").strip()

    busy_ids = set(
        GuaranteeReleaseRequest.objects.filter(
            guarantee__client_id=client.pk,
            status__in=_OPEN_RELEASE_STATUSES,
        ).values_list("guarantee_id", flat=True)
    )
    guarantees_qs = (
        Guarantee.objects.filter(
            client_id=client.pk,
            status=Guarantee.Status.ACTIVE,
        )
        .exclude(id__in=busy_ids)
        .select_related("application", "agency")
        .order_by("-created_at")
    )

    apps = (
        CreditApplication.objects.filter(client_id=client.pk)
        .select_related("product")
        .order_by("-created_at")
    )
    loans_by_app = {
        loan.application_id: loan
        for loan in Loan.objects.filter(application__client_id=client.pk)
    }

    credits = []
    client_ids = _client_cbs_ids(client)
    for app in apps:
        loan = loans_by_app.get(app.pk)
        currency = getattr(app, "currency", None) or tenant_currency(tid)
        loan_ref = _loan_cbs_ref(loan)
        if not loan_ref and getattr(app, "client_account_number", None):
            # Dernier recours legacy (hors Perfect refDemande).
            loan_ref = (app.client_account_number or "").strip()

        cbs_info = {
            "cbs_loan_reference": loan_ref,
            "cbs_settled": None,
            "cbs_outstanding": None,
            "cbs_currency": currency,
            "cbs_error": None,
            "cbs_status_label": "Non vérifié",
        }
        if loan_ref:
            try:
                cbs = get_loan_status(
                    tid, loan_ref, currency=currency, **client_ids
                )
                cbs_info.update(
                    {
                        "cbs_settled": cbs["settled"],
                        "cbs_outstanding": str(cbs["outstanding"]),
                        "cbs_currency": cbs["currency"],
                        "cbs_status_label": "Soldé" if cbs["settled"] else "Actif",
                    }
                )
            except CoreBankingError as exc:
                cbs_info["cbs_error"] = str(exc)
                cbs_info["cbs_status_label"] = "Erreur CBS"

        credits.append(
            {
                "application_id": str(app.pk),
                "application_reference": app.reference,
                "application_status": app.status,
                "application_status_display": app.get_status_display(),
                "product_label": getattr(app.product, "label", "") if app.product_id else "",
                "amount": str(reference_amount(app) or 0),
                "currency": currency,
                "loan_id": str(loan.pk) if loan else None,
                "loan_status": loan.status if loan else None,
                "loan_status_display": loan.get_status_display() if loan else None,
                "disbursed_at": loan.disbursed_at.isoformat() if loan and loan.disbursed_at else None,
                **cbs_info,
            }
        )

    return {
        "client_id": str(client.pk),
        "client_display": getattr(client, "display_name", str(client)),
        "cbs_client_id": cbs_id,
        "guarantees": GuaranteeSerializer(guarantees_qs, many=True).data,
        "credits": credits,
    }


def _notify_collection_on_release_complete(request: GuaranteeReleaseRequest):
    if not request.application_id:
        return
    try:
        from apps.collections.models import CollectionCase
        from apps.collections.services import set_next_action
    except Exception:  # noqa: BLE001
        return
    case = (
        CollectionCase.objects.filter(loan__application_id=request.application_id)
        .exclude(stage=CollectionCase.Stage.CLOSED)
        .order_by("-opened_at")
        .first()
    )
    if case is None:
        return
    note = f"Main levée {request.reference} clôturée — garantie libérée."[:255]
    try:
        set_next_action(
            case,
            action_date=timezone.localdate(),
            action_type="RELEASE",
            note=note,
        )
    except Exception:  # noqa: BLE001
        pass


@transaction.atomic
def complete_release_request(request: GuaranteeReleaseRequest):
    """
    Re-vérifie CBS, exige l'acte signé (règle A), puis libère la garantie.
    """
    if not request.has_signed_acte():
        raise ProcessError(
            "L'acte de main levée signé doit être déposé avant la clôture."
        )

    guarantee = request.guarantee
    currency = request.cbs_currency or tenant_currency(request.tenant_id)
    client_ids = _client_cbs_ids(getattr(guarantee, "client", None))
    if request.cbs_client_id:
        client_ids["code_adherent"] = request.cbs_client_id.strip()
    try:
        cbs = assert_loan_settled(
            request.tenant_id,
            request.cbs_loan_reference,
            currency=currency,
            **client_ids,
        )
    except CoreBankingError as exc:
        request.status = GuaranteeReleaseRequest.Status.BLOCKED
        request.cbs_settled = False
        request.cbs_checked_at = timezone.now()
        request.comment = (
            (request.comment + "\n" if request.comment else "")
            + f"[CBS] {exc}"
        )
        request.save()
        return request

    request.cbs_settled = cbs["settled"]
    request.cbs_outstanding = cbs["outstanding"]
    request.cbs_currency = cbs["currency"]
    request.cbs_checked_at = timezone.now()
    request.cbs_raw = cbs.get("raw") or {}

    GuaranteeMovement.objects.create(
        tenant_id=guarantee.tenant_id,
        guarantee=guarantee,
        movement_type=GuaranteeMovement.MovementType.RELEASE,
        movement_date=timezone.now().date(),
        value=guarantee.current_value,
        comment=f"Main levée {request.reference}",
    )
    guarantee.status = Guarantee.Status.RELEASED
    guarantee.save(update_fields=["status", "updated_at"])

    request.status = GuaranteeReleaseRequest.Status.COMPLETED
    request.completed_at = timezone.now()
    request.save()
    _notify_collection_on_release_complete(request)
    return request


def _to_decimal_or_none(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise ProcessError(f"Valeur numérique invalide : {value}") from exc


def _busy_guarantee_ids_for_dation(*, client_id, exclude_dation_id=None):
    qs = DationAsset.objects.filter(
        dation__client_id=client_id,
        dation__status__in=_OPEN_DATION_STATUSES,
        guarantee__isnull=False,
    )
    if exclude_dation_id:
        qs = qs.exclude(dation_id=exclude_dation_id)
    return set(str(x) for x in qs.values_list("guarantee_id", flat=True))


def _assert_dation_editable(request: DationRequest):
    if request.status not in (
        DationRequest.Status.DRAFT,
        DationRequest.Status.RETURNED,
    ):
        raise ProcessError(
            "Seuls les dossiers brouillon ou retournés peuvent être modifiés."
        )


def _refresh_dation_summary(request: DationRequest):
    lines = list(request.assets.all())
    summary = " | ".join(
        (line.description or "Bien").strip() for line in lines
    )[:2000]
    request.asset_description = summary
    request.apply_settlement_snapshot(persist=False)
    request.save(
        update_fields=[
            "asset_description",
            "asset_value",
            "fees_client_total",
            "fees_institution_total",
            "claim_to_cover",
            "residual_balance",
            "surplus_amount",
            "updated_at",
        ]
    )


def _build_dation_asset_lines(
    *,
    client,
    guarantee_ids,
    additional_assets,
    legacy_description="",
    legacy_value=None,
    exclude_dation_id=None,
):
    """Construit la liste normalisée des biens du dossier."""
    lines = []
    seen = set()
    busy = _busy_guarantee_ids_for_dation(
        client_id=client.pk, exclude_dation_id=exclude_dation_id
    )

    for raw_id in guarantee_ids:
        gid = str(raw_id).strip()
        if not gid or gid in seen:
            continue
        seen.add(gid)
        try:
            guarantee = Guarantee.objects.get(pk=gid, client_id=client.pk)
        except Guarantee.DoesNotExist as exc:
            raise ProcessError(
                f"Garantie introuvable pour ce client : {gid}."
            ) from exc
        if guarantee.status != Guarantee.Status.ACTIVE:
            raise ProcessError(
                f"La garantie {guarantee.reference or gid} n'est pas active."
            )
        if gid in busy:
            raise ProcessError(
                f"La garantie {guarantee.reference or gid} est déjà "
                "engagée dans une autre dation ouverte."
            )
        value = (
            guarantee.current_value
            or guarantee.value_to_consider
            or guarantee.expertise_value
        )
        label = guarantee.get_guarantee_type_display()
        ref = guarantee.reference or str(guarantee.pk)[:8]
        desc = f"{label} — {ref}"
        if guarantee.description:
            desc = f"{desc} : {guarantee.description}"
        asset_type = DationAsset.AssetType.OTHER
        gtype = guarantee.guarantee_type
        if gtype == Guarantee.GuaranteeType.MORTGAGE:
            asset_type = DationAsset.AssetType.REAL_ESTATE
        elif gtype == Guarantee.GuaranteeType.PLEDGE:
            asset_type = DationAsset.AssetType.VEHICLE
        elif gtype == Guarantee.GuaranteeType.FINANCIAL:
            asset_type = DationAsset.AssetType.FINANCIAL
        lines.append(
            {
                "source": DationAsset.Source.EXISTING_GUARANTEE,
                "guarantee": guarantee,
                "description": desc,
                "value": value,
                "asset_type": asset_type,
                "notes": "",
            }
        )

    for item in additional_assets:
        if not isinstance(item, dict):
            raise ProcessError("Chaque bien additionnel doit être un objet.")
        desc = (item.get("description") or "").strip()
        value = _to_decimal_or_none(item.get("value"))
        if not desc and value is None:
            continue
        if not desc:
            raise ProcessError(
                "La description est obligatoire pour un bien additionnel."
            )
        asset_type = (item.get("asset_type") or DationAsset.AssetType.OTHER).strip()
        if asset_type not in DationAsset.AssetType.values:
            asset_type = DationAsset.AssetType.OTHER
        lines.append(
            {
                "source": DationAsset.Source.ADDITIONAL,
                "guarantee": None,
                "description": desc,
                "value": value,
                "asset_type": asset_type,
                "notes": (item.get("notes") or "").strip(),
            }
        )

    # Compatibilité ancien formulaire (un seul bien libre).
    if not lines and (legacy_description or legacy_value not in (None, "")):
        lines.append(
            {
                "source": DationAsset.Source.ADDITIONAL,
                "guarantee": None,
                "description": (legacy_description or "").strip() or "Bien cédé",
                "value": _to_decimal_or_none(legacy_value),
                "asset_type": DationAsset.AssetType.OTHER,
                "notes": "",
            }
        )
    return lines


def _create_fee_lines(request: DationRequest, fees):
    if not fees:
        return
    if not isinstance(fees, list):
        raise ProcessError("fees doit être une liste.")
    for item in fees:
        if not isinstance(item, dict):
            raise ProcessError("Chaque frais doit être un objet.")
        amount = _to_decimal_or_none(item.get("amount"))
        if amount is None:
            raise ProcessError("Le montant du frais est obligatoire.")
        fee_type = (item.get("fee_type") or DationFee.FeeType.OTHER).strip()
        if fee_type not in DationFee.FeeType.values:
            fee_type = DationFee.FeeType.OTHER
        payer = (item.get("payer") or DationFee.Payer.CLIENT).strip()
        if payer not in DationFee.Payer.values:
            payer = DationFee.Payer.CLIENT
        label = (item.get("label") or "").strip()
        if not label:
            label = dict(DationFee.FeeType.choices).get(fee_type, "Frais")
        fee_date = item.get("fee_date") or None
        asset = None
        asset_id = item.get("asset") or item.get("asset_id")
        if asset_id:
            try:
                asset = request.assets.get(pk=asset_id)
            except DationAsset.DoesNotExist as exc:
                raise ProcessError("Bien lié au frais introuvable.") from exc
        DationFee.objects.create(
            tenant_id=request.tenant_id,
            dation=request,
            asset=asset,
            fee_type=fee_type,
            label=label,
            amount=amount,
            payer=payer,
            fee_date=fee_date,
            recoverable=bool(item.get("recoverable", True)),
            notes=(item.get("notes") or "").strip(),
        )


@transaction.atomic
def initiate_dation_request(
    *,
    client,
    user,
    asset_description="",
    asset_value=None,
    application=None,
    comment="",
    cbs_client_id="",
    guarantee_ids=None,
    additional_assets=None,
    fees=None,
    as_draft=True,
    require_full_coverage=False,
    settlement_notes="",
):
    """
    Vérifie l'encours CBS client puis crée le dossier de dation.

    Par défaut en brouillon (`as_draft=True`) pour permettre l'ajout de
    pièces / frais avant soumission au circuit. Avec `as_draft=False`,
    démarre immédiatement le circuit DATION.
    """
    cbs_id = (cbs_client_id or getattr(client, "cbs_client_id", "") or "").strip()
    currency = tenant_currency(client.tenant_id)
    if application is not None and getattr(application, "currency", None):
        currency = application.currency

    try:
        cbs = assert_client_outstanding_for_dation(
            client.tenant_id, cbs_id, currency=currency
        )
    except CoreBankingError as exc:
        raise ProcessError(str(exc)) from exc

    lines = _build_dation_asset_lines(
        client=client,
        guarantee_ids=guarantee_ids or [],
        additional_assets=additional_assets or [],
        legacy_description=asset_description or "",
        legacy_value=asset_value,
    )
    if not lines:
        raise ProcessError(
            "Ajoutez au moins une garantie existante ou un bien additionnel "
            "au dossier de dation."
        )

    total_value = sum((line["value"] or Decimal("0")) for line in lines)
    summary = " | ".join(
        (line["description"] or "Bien").strip() for line in lines
    )[:2000]

    req = DationRequest(
        client=client,
        application=application,
        agency_id=getattr(application, "agency_id", None) if application else None,
        cbs_client_id=cbs_id,
        cbs_total_outstanding=cbs["total_outstanding"],
        cbs_currency=cbs["currency"],
        cbs_checked_at=timezone.now(),
        cbs_raw=cbs.get("raw") or {},
        asset_description=summary,
        asset_value=total_value,
        require_full_coverage=bool(require_full_coverage),
        settlement_notes=settlement_notes or "",
        status=DationRequest.Status.DRAFT,
        comment=comment or "",
        created_by=user,
        updated_by=user,
    )
    req.reference = _next_reference("DAT", DationRequest, client.tenant_id)
    req.save()

    for line in lines:
        DationAsset.objects.create(
            tenant_id=client.tenant_id,
            dation=req,
            source=line["source"],
            guarantee=line.get("guarantee"),
            description=line["description"],
            value=line["value"],
            asset_type=line.get("asset_type") or DationAsset.AssetType.OTHER,
            notes=line.get("notes") or "",
        )

    _create_fee_lines(req, fees or [])
    req.apply_settlement_snapshot(persist=True)

    if as_draft:
        return req
    return submit_dation_request(req, user=user)


@transaction.atomic
def add_dation_asset(
    request: DationRequest,
    *,
    guarantee_id=None,
    description="",
    value=None,
    asset_type=None,
    notes="",
    user=None,
):
    _assert_dation_editable(request)
    if guarantee_id:
        lines = _build_dation_asset_lines(
            client=request.client,
            guarantee_ids=[guarantee_id],
            additional_assets=[],
            exclude_dation_id=request.pk,
        )
        line = lines[0]
    else:
        lines = _build_dation_asset_lines(
            client=request.client,
            guarantee_ids=[],
            additional_assets=[
                {
                    "description": description,
                    "value": value,
                    "asset_type": asset_type,
                    "notes": notes,
                }
            ],
            exclude_dation_id=request.pk,
        )
        if not lines:
            raise ProcessError("Bien additionnel invalide.")
        line = lines[0]
    asset = DationAsset.objects.create(
        tenant_id=request.tenant_id,
        dation=request,
        source=line["source"],
        guarantee=line.get("guarantee"),
        description=line["description"],
        value=line["value"],
        asset_type=line.get("asset_type") or DationAsset.AssetType.OTHER,
        notes=line.get("notes") or "",
    )
    if user is not None:
        request.updated_by = user
        request.save(update_fields=["updated_by", "updated_at"])
    _refresh_dation_summary(request)
    return asset


@transaction.atomic
def remove_dation_asset(request: DationRequest, asset_id, *, user=None):
    _assert_dation_editable(request)
    try:
        asset = request.assets.get(pk=asset_id)
    except DationAsset.DoesNotExist as exc:
        raise ProcessError("Bien introuvable sur ce dossier.") from exc
    asset.delete()
    if user is not None:
        request.updated_by = user
        request.save(update_fields=["updated_by", "updated_at"])
    _refresh_dation_summary(request)
    return request


@transaction.atomic
def add_dation_fee(request: DationRequest, data: dict, *, user=None):
    _assert_dation_editable(request)
    _create_fee_lines(request, [data])
    if user is not None:
        request.updated_by = user
        request.save(update_fields=["updated_by", "updated_at"])
    _refresh_dation_summary(request)
    return request.fees.order_by("-created_at").first()


@transaction.atomic
def remove_dation_fee(request: DationRequest, fee_id, *, user=None):
    _assert_dation_editable(request)
    try:
        fee = request.fees.get(pk=fee_id)
    except DationFee.DoesNotExist as exc:
        raise ProcessError("Frais introuvable sur ce dossier.") from exc
    fee.delete()
    if user is not None:
        request.updated_by = user
        request.save(update_fields=["updated_by", "updated_at"])
    _refresh_dation_summary(request)
    return request


@transaction.atomic
def update_dation_request(request: DationRequest, *, user=None, **fields):
    _assert_dation_editable(request)
    if "comment" in fields and fields["comment"] is not None:
        request.comment = fields["comment"]
    if "settlement_notes" in fields and fields["settlement_notes"] is not None:
        request.settlement_notes = fields["settlement_notes"]
    if "require_full_coverage" in fields and fields["require_full_coverage"] is not None:
        request.require_full_coverage = bool(fields["require_full_coverage"])
    if "application" in fields:
        request.application = fields["application"]
        request.agency_id = getattr(fields["application"], "agency_id", None) if fields["application"] else request.agency_id
    if user is not None:
        request.updated_by = user
    request.save()
    _refresh_dation_summary(request)
    return request


@transaction.atomic
def refresh_dation_cbs(request: DationRequest, *, user=None):
    if request.status in (
        DationRequest.Status.COMPLETED,
        DationRequest.Status.CANCELLED,
        DationRequest.Status.REJECTED,
    ):
        raise ProcessError("Impossible de rafraîchir le CBS sur ce statut.")
    currency = request.cbs_currency or tenant_currency(request.tenant_id)
    try:
        cbs = assert_client_outstanding_for_dation(
            request.tenant_id,
            request.cbs_client_id,
            currency=currency,
        )
    except CoreBankingError as exc:
        raise ProcessError(str(exc)) from exc
    request.cbs_total_outstanding = cbs["total_outstanding"]
    request.cbs_currency = cbs["currency"]
    request.cbs_checked_at = timezone.now()
    request.cbs_raw = cbs.get("raw") or {}
    if user is not None:
        request.updated_by = user
    request.save()
    _refresh_dation_summary(request)
    return request


@transaction.atomic
def submit_dation_request(request: DationRequest, *, user=None):
    """Passe le brouillon / retourné en validation et démarre le circuit."""
    if request.status not in (
        DationRequest.Status.DRAFT,
        DationRequest.Status.RETURNED,
    ):
        raise ProcessError(
            "Seuls les dossiers brouillon ou retournés peuvent être soumis."
        )
    if not request.assets.exists():
        raise ProcessError("Ajoutez au moins un bien avant de soumettre.")

    refresh_dation_cbs(request, user=user)
    settlement = request.apply_settlement_snapshot(persist=True)
    if request.require_full_coverage and not settlement["covers_claim"]:
        raise ProcessError(
            "Couverture insuffisante : la créance à couvrir n'est pas "
            "entièrement couverte par les biens."
        )

    amount = settlement["assets_total"] or settlement["claim_to_cover"] or 0
    try:
        start_workflow(
            request,
            amount=amount,
            risk_level=None,
            target_type=WorkflowDefinition.TargetType.DATION,
        )
    except WorkflowError as exc:
        raise ProcessError(str(exc)) from exc

    request.status = DationRequest.Status.IN_APPROVAL
    if user is not None:
        request.updated_by = user
    request.save(update_fields=["status", "updated_by", "updated_at"] if user else ["status", "updated_at"])
    return request


@transaction.atomic
def cancel_dation_request(request: DationRequest, *, user=None, comment=""):
    if request.status not in (
        DationRequest.Status.DRAFT,
        DationRequest.Status.RETURNED,
        DationRequest.Status.IN_APPROVAL,
        DationRequest.Status.BLOCKED,
    ):
        raise ProcessError("Ce dossier ne peut plus être annulé.")
    from apps.workflow.services import cancel_active_workflows_for_target

    cancel_active_workflows_for_target(request)
    request.status = DationRequest.Status.CANCELLED
    if comment:
        request.comment = (
            (request.comment + "\n" if request.comment else "") + comment
        )
    if user is not None:
        request.updated_by = user
    request.save()
    return request


def preview_dation_cbs(*, tenant_id, cbs_client_id, currency=None) -> dict:
    """Lecture encours CBS (créance) sans créer de demande."""
    currency = (currency or "").strip().upper() or tenant_currency(tenant_id)
    try:
        return get_client_outstanding(
            tenant_id, cbs_client_id, currency=currency
        )
    except CoreBankingError as exc:
        raise ProcessError(str(exc)) from exc


def ensure_dation_document_categories(tenant) -> int:
    """Crée les catégories GED dation si absentes."""
    from apps.documents.category_seed import ensure_document_categories

    return ensure_document_categories(
        tenant,
        (
            ("DAT_ACTE", "Acte de dation"),
            ("DAT_PHOTO", "Photo du bien"),
            ("DAT_EXPERTISE", "Rapport d'expertise"),
            ("DAT_TITRE", "Titre / carte grise / justificatif"),
            ("DAT_FACTURE", "Facture / quittance de frais"),
            ("DAT_OTHER", "Autre pièce dation"),
        ),
    )


def dation_documents(request: DationRequest):
    from django.contrib.contenttypes.models import ContentType

    from apps.documents.models import Document

    ct_req = ContentType.objects.get_for_model(DationRequest)
    ct_asset = ContentType.objects.get_for_model(DationAsset)
    asset_ids = list(request.assets.values_list("id", flat=True))
    from django.db.models import Q

    q = Q(content_type=ct_req, object_id=request.id)
    if asset_ids:
        q |= Q(content_type=ct_asset, object_id__in=asset_ids)
    return Document.objects.filter(q).select_related("category", "uploaded_by")


def _notify_collection_on_dation_complete(request: DationRequest):
    """Note légère sur le dossier de recouvrement lié, si présent."""
    if not request.application_id:
        return
    try:
        from apps.collections.models import CollectionCase
        from apps.collections.services import set_next_action
    except Exception:  # noqa: BLE001
        return
    case = (
        CollectionCase.objects.filter(
            loan__application_id=request.application_id,
        )
        .exclude(stage=CollectionCase.Stage.CLOSED)
        .order_by("-opened_at")
        .first()
    )
    if case is None:
        return
    note = (
        f"Dation {request.reference} clôturée — "
        f"couverture={request.covers_claim()}, "
        f"résiduel={request.residual_balance}"
    )[:255]
    try:
        set_next_action(
            case,
            action_date=timezone.localdate(),
            action_type="DATION",
            note=note,
        )
    except Exception:  # noqa: BLE001
        pass


@transaction.atomic
def complete_dation_request(request: DationRequest):
    """
    Re-vérifie l'encours CBS, réalise les garanties sources, puis enregistre
    la garantie type DATION résultante.
    """
    currency = request.cbs_currency or tenant_currency(request.tenant_id)
    try:
        cbs = assert_client_outstanding_for_dation(
            request.tenant_id,
            request.cbs_client_id,
            currency=currency,
        )
    except CoreBankingError as exc:
        request.status = DationRequest.Status.BLOCKED
        request.cbs_checked_at = timezone.now()
        request.comment = (
            (request.comment + "\n" if request.comment else "")
            + f"[CBS] {exc}"
        )
        request.save()
        return request

    request.cbs_total_outstanding = cbs["total_outstanding"]
    request.cbs_currency = cbs["currency"]
    request.cbs_checked_at = timezone.now()
    request.cbs_raw = cbs.get("raw") or {}
    settlement = request.apply_settlement_snapshot(persist=False)

    # Sort des garanties sources : réalisation + mouvement.
    for asset in request.assets.select_related("guarantee"):
        guarantee = asset.guarantee
        if guarantee is None or guarantee.status != Guarantee.Status.ACTIVE:
            continue
        GuaranteeMovement.objects.create(
            tenant_id=guarantee.tenant_id,
            guarantee=guarantee,
            movement_type=GuaranteeMovement.MovementType.REALIZATION,
            movement_date=timezone.now().date(),
            value=asset.value or guarantee.current_value,
            comment=f"Réalisée par dation {request.reference}",
        )
        guarantee.status = Guarantee.Status.REALIZED
        guarantee.save(update_fields=["status", "updated_at"])

    value = settlement["assets_total"] or cbs["total_outstanding"]
    description = request.asset_description or (
        f"Dation en paiement {request.reference}"
    )
    if settlement["residual_balance"]:
        description = (
            f"{description} | Solde résiduel : {settlement['residual_balance']}"
        )
    guarantee = Guarantee.objects.create(
        client=request.client,
        application=request.application,
        agency_id=request.agency_id,
        guarantee_type=Guarantee.GuaranteeType.DATION,
        description=description,
        expertise_value=value or 0,
        current_value=value or 0,
        status=Guarantee.Status.ACTIVE,
        created_by=request.created_by,
        updated_by=request.updated_by,
    )
    guarantee.reference = _next_reference("GAR-DAT", Guarantee, request.tenant_id)
    guarantee.save(update_fields=["reference", "updated_at"])

    request.resulting_guarantee = guarantee
    request.status = DationRequest.Status.COMPLETED
    request.completed_at = timezone.now()
    request.save()
    _notify_collection_on_dation_complete(request)
    return request
