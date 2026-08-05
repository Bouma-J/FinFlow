"""Moteurs de rendu : injecte les données dans les modèles en préservant la forme.

- ``.docx`` : via ``docxtpl`` (Jinja2 dans Word) — mise en page conservée.
- ``.xlsx`` : via ``openpyxl`` — substitution des balises ``{{ var }}`` dans les
  cellules, format conservé.
"""
from __future__ import annotations

import io
import re

from django.core.files.base import ContentFile

_TOKEN_RE = re.compile(r"{{\s*([a-zA-Z0-9_]+)\s*}}")


class ContractRenderError(Exception):
    """Erreur de rendu d'un contrat (modèle invalide, balise inconnue…)."""


def render_docx(template_file, context: dict) -> bytes:
    """Rend un modèle Word et retourne les octets du document produit."""
    try:
        from docxtpl import DocxTemplate
    except ImportError as exc:  # pragma: no cover
        raise ContractRenderError(
            "La dépendance 'docxtpl' n'est pas installée."
        ) from exc

    template_file.open("rb")
    try:
        buffer = io.BytesIO(template_file.read())
    finally:
        template_file.close()

    try:
        doc = DocxTemplate(buffer)
        doc.render(context)
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception as exc:
        raise ContractRenderError(
            f"Impossible de générer le contrat Word : {exc}"
        ) from exc


def render_xlsx(template_file, context: dict) -> bytes:
    """Rend un modèle Excel (substitution des balises scalaires)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ContractRenderError(
            "La dépendance 'openpyxl' n'est pas installée."
        ) from exc

    scalars = {
        k: ("" if v is None else str(v))
        for k, v in context.items()
        if not isinstance(v, (list, dict))
    }

    def substitute(match):
        key = match.group(1)
        return scalars.get(key, match.group(0))

    template_file.open("rb")
    try:
        buffer = io.BytesIO(template_file.read())
    finally:
        template_file.close()

    try:
        wb = load_workbook(buffer)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "{{" in cell.value:
                        cell.value = _TOKEN_RE.sub(substitute, cell.value)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except Exception as exc:
        raise ContractRenderError(
            f"Impossible de générer le contrat Excel : {exc}"
        ) from exc


def render_template(template, context: dict) -> ContentFile:
    """Rend un ``ContractTemplate`` et retourne un ``ContentFile`` nommé."""
    from .models import ContractTemplate

    name = (template.file.name or "").lower()
    is_xlsx = template.engine == ContractTemplate.Engine.XLSX or name.endswith(".xlsx")

    if is_xlsx:
        data = render_xlsx(template.file, context)
        ext = "xlsx"
    else:
        data = render_docx(template.file, context)
        ext = "docx"

    ref = context.get("dossier_reference") or "dossier"
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", template.name).strip("_") or "contrat"
    filename = f"{slug}_{ref}.{ext}"
    return ContentFile(data, name=filename)
